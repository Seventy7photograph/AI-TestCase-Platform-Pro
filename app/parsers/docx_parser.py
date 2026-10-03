"""Word (.docx) 解析器。

相比只读段落，这里额外把表格转成「表头: 值」的形式，
因为需求文档里的字段约束（范围/长度/枚举）绝大多数写在表格中，
保留表格结构能显著提升后续字段约束抽取的准确率。
"""
from __future__ import annotations

import io

from app.core.exceptions import DocumentParseError
from app.parsers.base import DocumentParser, ParsedDocument

try:  # pragma: no cover - 取决于运行环境是否安装 python-docx
    from docx import Document as _DocxDocument

    DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover
    _DocxDocument = None  # type: ignore[assignment]
    DOCX_AVAILABLE = False


class DocxParser(DocumentParser):
    name = "docx"
    label = "Word 解析器"
    extensions = ("docx",)

    def parse(self, data: bytes, *, filename: str = "") -> ParsedDocument:
        if not DOCX_AVAILABLE:  # pragma: no cover
            raise DocumentParseError(
                "解析 .docx 需要安装 python-docx，请执行：pip install python-docx",
                detail={"missing_dependency": "python-docx"},
            )

        try:
            document = _DocxDocument(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001 - 第三方库异常类型不稳定，统一兜底
            raise DocumentParseError(
                f"Word 文档打开失败，文件可能已损坏或受密码保护：{exc}",
                detail={"filename": filename, "error": str(exc)},
            ) from exc

        lines: list[str] = []
        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()
            if not text:
                continue
            lines.append(f"{self._heading_prefix(paragraph)}{text}")

        warnings: list[str] = []
        table_count = 0
        for table in document.tables:
            table_count += 1
            lines.append("")
            lines.append(f"[表格 {table_count}]")
            lines.extend(self._render_table(table))

        if table_count:
            warnings.append(f"文档包含 {table_count} 张表格，已按「表头: 值」形式平铺，便于抽取字段约束。")

        return self._finalize(
            "\n".join(lines),
            extension="docx",
            table_count=table_count,
            metadata={"filename": filename, "paragraph_count": len(document.paragraphs)},
            warnings=warnings,
        )

    @staticmethod
    def _heading_prefix(paragraph) -> str:
        """把 Word 标题样式还原成 Markdown 的 #，供规则解析器识别章节。"""
        try:
            style_name = (paragraph.style.name or "").lower()
        except Exception:  # noqa: BLE001
            return ""
        if "heading" in style_name or "标题" in style_name:
            level = "".join(ch for ch in style_name if ch.isdigit())
            return "#" * (int(level) if level else 1) + " "
        return ""

    @staticmethod
    def _render_table(table) -> list[str]:
        rows = list(table.rows)
        if not rows:
            return []
        header_cells = [cell.text.strip().replace("\n", " ") for cell in rows[0].cells]
        rendered: list[str] = ["| " + " | ".join(header_cells) + " |"]
        for row in rows[1:]:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            pairs = [
                f"{header_cells[index] or f'列{index + 1}'}: {value}"
                for index, value in enumerate(values)
                if value
            ]
            if pairs:
                rendered.append("| " + " | ".join(pairs) + " |")
        return rendered